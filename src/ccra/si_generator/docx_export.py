from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from .models import SIDocument, TableData

FONT = "Times New Roman"
SIZE = Pt(9)


def _set_run(run, *, bold: bool | None = None, superscript: bool = False) -> None:
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = SIZE
    if bold is not None:
        run.bold = bold
    run.font.superscript = superscript


def _set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(str(text))
    _set_run(r, bold=bold)


def _set_cell_borders(cell, top: bool, bottom: bool) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for side_name, enabled in (("top", top), ("bottom", bottom)):
        side = OxmlElement(f"w:{side_name}")
        side.set(qn("w:val"), "single" if enabled else "nil")
        side.set(qn("w:sz"), "8")
        side.set(qn("w:color"), "000000")
        borders.append(side)
    for side_name in ("left", "right", "insideH", "insideV"):
        side = OxmlElement(f"w:{side_name}")
        side.set(qn("w:val"), "nil")
        borders.append(side)


def _add_table(doc: Document, data: TableData) -> None:
    rows = data.normalized_rows()
    table = doc.add_table(rows=1 + len(rows), cols=len(data.headers))
    table.autofit = True
    for col, header in enumerate(data.headers):
        _set_cell_text(table.cell(0, col), header, bold=True)
    for row_index, values in enumerate(rows, start=1):
        for col, value in enumerate(values):
            _set_cell_text(table.cell(row_index, col), value, bold=(col == 0))
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            _set_cell_borders(cell, top=(row_index == 0), bottom=(row_index in {0, len(table.rows) - 1}))


def _caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run("Table S. ")
    _set_run(r, bold=True)
    r = p.add_run(text)
    _set_run(r, bold=False)


def _footnote(doc: Document, mark: str, text: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(mark)
    _set_run(r, superscript=True)
    r = p.add_run(f" {text}")
    _set_run(r)


def export_docx(document: SIDocument, target: Path) -> Path:
    target = Path(target)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Pt(54)

    mechanism = document.mechanism_name.strip() or document.name.strip() or "reaction mechanism"

    _add_table(doc, document.absolute)
    _caption(doc, f"Absolute energy values of {mechanism} in Kcal/Mol.")

    _add_table(doc, document.relative)
    _caption(doc, f"Relative energy values of {mechanism} in Kcal/Mol.")
    _footnote(doc, "a", "Def2-TZVPP electronic energy = sum of the two previous columns.")
    _footnote(doc, "b", f"T = {document.temperature_k:g} K.")
    _footnote(doc, "c", "Gibb’s free energy = sum of the previous six columns.")
    doc.add_paragraph()

    _add_table(doc, document.spin_density)
    _caption(doc, f"Mulliken Spin Density Distributions of {mechanism}.")

    _add_table(doc, document.geometry)
    _caption(doc, f"Selected Geometries of {mechanism} in Å and °.")

    if document.notes.strip():
        p = doc.add_paragraph()
        r = p.add_run(document.notes.strip())
        _set_run(r)

    for p in doc.paragraphs:
        p.paragraph_format.space_after = Pt(2)
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(target)
    return target
